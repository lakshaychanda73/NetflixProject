#!/usr/bin/env python3
"""Merge researcher JSONL files -> deduplicated Mumbai/MMR builders outreach workbook."""
import glob
import json
import os
import re
import sys
from collections import OrderedDict, defaultdict

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

NOTES = "/home/user/NetflixProject/research_notes/Mumbai MMR builders outreach database"
OUT = "/home/user/NetflixProject/reports/Mumbai MMR builders outreach database.xlsx"
SCRATCH = os.path.dirname(os.path.abspath(__file__))
ALIAS_FILE = os.path.join(SCRATCH, "aliases.json")        # {"merge": [[name, name, ...], ...], "keep_apart": [[a, b], ...], "level_override": {name: level}}
CHECKS_GLOB = os.path.join(NOTES, "qa_contact_checks*.jsonl")
EXAMPLES_FILE = os.path.join(SCRATCH, "level_examples.json")
AS_OF = "September 2026"
BUILD_DATE = "2026-09-29"

NA = "Not publicly available"
NF = "Not found"
LEVEL_ORDER = {"L1": 1, "L2": 2, "L3": 3, "L4": 4, "L5": 5, "Connector": 6}
ROLE_ORDER = ["Promoter/Founder", "Chairman", "MD/CEO", "Director/Board", "Business Development",
              "Land Acquisition", "Redevelopment", "Sales", "Channel Sales", "Strategy", "Marketing",
              "CFO/Finance", "Investor Relations", "Company Secretary", "Liaison/Approvals",
              "Association Office Bearer", "Other", "Company contact"]
ROLE_RANK = {r: i for i, r in enumerate(ROLE_ORDER)}

STOP = {"ltd", "limited", "pvt", "private", "llp", "the", "co", "company", "inc", "formerly", "ex", "india", "and", "of"}
GENERIC = {"group", "groups", "developers", "developer", "realty", "realtors", "realtor", "builders", "builder", "infra",
           "infrastructure", "infrastructures", "infraconstruction", "properties", "property", "estates", "estate",
           "constructions", "construction", "housing", "lifespaces", "lifespace", "spaces", "projects", "project",
           "ventures", "homes", "corp", "corporation", "landmarks", "landmark", "enterprises", "promoters",
           "associates", "buildcon", "lifestyle", "lifestyles", "world", "real", "developments", "development",
           "shelters", "shelter", "creators", "habitat", "habitats", "structures", "superstructures", "reality", "mumbai"}

NA_RE = re.compile(r"^\s*(|n/?a|na|none|null|nil|-+|unknown|unavailable|not (publicly )?(available|found|listed|disclosed|published)"
                   r"|not public(ly)?( available| listed| disclosed)?)\s*\.?\s*$", re.I)
GENERIC_EMAIL = re.compile(r"^(info|sales|enquir|inquir|contact|customer|care|support|hello|admin|office|investor|ir|cs|secretar|"
                           r"compliance|grievance|marketing|media|pr|cp|channel|partner|booking|crm|hr|career|jobs|legal|"
                           r"corporate|mail|welcome|reach|connect|feedback|enquiry|query|queries|help|services|accounts|"
                           r"shareholder|redevelop|land|bd|business)", re.I)
URL_RE = re.compile(r"https?://[^\s,;\]\)\"']+")


# ---------------------------------------------------------------- helpers
def s(v):
    if v is None:
        return ""
    if isinstance(v, (list, tuple)):
        return "; ".join(s(x) for x in v if s(x))
    if isinstance(v, dict):
        return "; ".join(f"{k}: {s(x)}" for k, x in v.items() if s(x))
    return str(v).strip()


def is_na(v):
    t = s(v)
    low = t.lower()
    return (not t) or bool(NA_RE.match(t)) or low.startswith("not ") or low.startswith("unverified") or low.startswith("none")


NV = "Not verified yet"
UNVERIFIED_HINT = re.compile(r"(?i)(not verified|unverified|not seen|not captured|blocked|unreachable|could not|couldn.t|not checked|re-?check|not researched)")


def clean_contact(v):
    t = s(v)
    if not t:
        return NA
    low = t.lower()
    if is_na(t) or low.startswith("not ") or low.startswith("unverified") or low.startswith("none"):
        return NV if UNVERIFIED_HINT.search(t) else NA
    return t


def clean_link(v):
    t = s(v)
    return NF if is_na(t) or not t.lower().startswith("http") else t


def src_list(v):
    if v is None:
        return []
    if isinstance(v, str):
        found = URL_RE.findall(v)
        return found if found else ([v.strip()] if v.strip() else [])
    out = []
    for x in v:
        out.extend(src_list(x) if not isinstance(x, str) else (URL_RE.findall(x) or ([x.strip()] if x.strip() else [])))
    return out


def uniq(seq):
    seen, out = set(), []
    for x in seq:
        k = x.strip().rstrip("/").lower() if isinstance(x, str) else x
        if k and k not in seen:
            seen.add(k)
            out.append(x.strip() if isinstance(x, str) else x)
    return out


def norm_tokens(name):
    t = name.lower().replace("&", " and ")
    t = re.sub(r"[^a-z0-9 ]", " ", t)
    return [w for w in t.split() if w not in STOP]


def strict_key(name):
    main = re.sub(r"\([^)]*\)", " ", name)
    return " ".join(norm_tokens(main))


def loose_keys(name):
    parts = [re.sub(r"\([^)]*\)", " ", name)] + re.findall(r"\(([^)]*)\)", name)
    keys = set()
    for p in parts:
        p = re.sub(r"(?i)\b(formerly|ex|earlier|erstwhile|now)\b.*", "", p)
        toks = [w for w in norm_tokens(p) if w not in GENERIC]
        if toks:
            keys.add(" ".join(toks))
    return keys


def domain(url):
    m = re.match(r"https?://([^/]+)", s(url).lower())
    if not m:
        return ""
    d = m.group(1)
    d = re.sub(r"^www\d?\.", "", d)
    if any(x in d for x in ["facebook.", "instagram.", "linkedin.", "justdial", "indiamart", "google.", "99acres", "magicbricks",
                            "housing.com", "maharera", "twitter.", "x.com", "youtube", "wixsite", "blogspot", "squarespace"]):
        return ""
    return d


def norm_person(n):
    t = re.sub(r"(?i)\b(mr|mrs|ms|dr|shri|smt|er|ca|adv|capt|col)\.?\s+", "", s(n))
    t = re.sub(r"[^a-z ]", " ", t.lower())
    return " ".join(t.split())


def phone_kind(v):
    t = s(v)
    if is_na(t):
        return None
    low = t.lower()
    if any(w in low for w in ["board", "office", "company", "landline", "reception", "corporate", "toll", "sales line", "helpline", "general"]):
        return "office"
    digits = re.sub(r"\D", "", t.split(";")[0].split("/")[0])
    if digits.startswith("91") and len(digits) == 12:
        digits = digits[2:]
    if digits.startswith("0") and len(digits) == 11:
        digits = digits[1:]
    if len(digits) == 10 and digits[0] in "6789":
        return "mobile"
    return "office"


def email_kind(v):
    t = s(v)
    if is_na(t):
        return None
    first = re.split(r"[;,/ ]", t)[0]
    local = first.split("@")[0]
    return "generic" if GENERIC_EMAIL.match(local) else "direct"


def contact_availability(mobile, email, basis=None):
    if basis == KB_LABEL and s(mobile) in ("", NV, NA) and s(email) in ("", NV, NA):
        return "Not verified yet – check website/MahaRERA"
    pk, ek = phone_kind(mobile), email_kind(email)
    direct_e, direct_p = ek == "direct", pk == "mobile"
    if direct_e and direct_p:
        return "Direct email + mobile public"
    if direct_e and pk:
        return "Direct email + office line public"
    if direct_e:
        return "Direct email public"
    if direct_p:
        return "Mobile public" + (" + role email" if ek else "")
    if pk or ek:
        return "Role/company email or line only"
    return "Not publicly available – use company line"


def basis_of(topic):
    return KB_LABEL if topic.startswith("kb_") else WEB_LABEL


KB_LABEL = "Model knowledge – unverified"
WEB_LABEL = "Web-sourced (search extracts, Sep 2026)"


def linkedin_search(name, company):
    from urllib.parse import quote
    brand = re.sub(r"\([^)]*\)", "", company).replace("Ltd", "").replace("Limited", "").replace("Pvt", "").strip()
    return "https://www.linkedin.com/search/results/people/?keywords=" + quote(f"{name} {brand}")


# ---------------------------------------------------------------- load
def load_jsonl(path):
    rows, bad = [], 0
    with open(path, encoding="utf-8") as f:
        txt = f.read()
    stripped = txt.strip()
    if stripped.startswith("["):
        try:
            return json.loads(stripped), 0
        except Exception:
            pass
    for line in txt.splitlines():
        line = line.strip().rstrip(",")
        if not line or line in ("[", "]"):
            continue
        try:
            obj = json.loads(line)
            if isinstance(obj, dict):
                rows.append(obj)
            elif isinstance(obj, list):
                rows.extend(o for o in obj if isinstance(o, dict))
        except Exception:
            bad += 1
    return rows, bad


def topic_of(path, suffix):
    return os.path.basename(path)[: -len(suffix)]


def main():
    aliases = {"merge": [], "keep_apart": [], "level_override": {}, "rename": {}}
    if os.path.exists(ALIAS_FILE):
        aliases.update(json.load(open(ALIAS_FILE)))
    report = defaultdict(list)

    companies = []   # dicts with _topic
    for p in sorted(glob.glob(os.path.join(NOTES, "*_companies.jsonl"))):
        rows, bad = load_jsonl(p)
        t = topic_of(p, "_companies.jsonl")
        if bad:
            report["bad_lines"].append(f"{os.path.basename(p)}: {bad} unparsable lines")
        for r in rows:
            name = s(r.get("company"))
            if not name:
                continue
            r["_topic"] = t
            r["company"] = aliases["rename"].get(name, name)
            companies.append(r)
    people = []
    for p in sorted(glob.glob(os.path.join(NOTES, "*_people.jsonl"))):
        rows, bad = load_jsonl(p)
        t = topic_of(p, "_people.jsonl")
        if bad:
            report["bad_lines"].append(f"{os.path.basename(p)}: {bad} unparsable lines")
        for r in rows:
            if not s(r.get("person")):
                continue
            r["_topic"] = t
            r["company"] = aliases["rename"].get(s(r.get("company")), s(r.get("company")))
            people.append(r)

    # ---------------- company clustering (union-find)
    names = uniq([c["company"] for c in companies] + [p["company"] for p in people if p["company"]])
    parent = {n: n for n in names}

    def find(x):
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x

    def union(a, b):
        if a in parent and b in parent:
            ra, rb = find(a), find(b)
            if ra != rb:
                parent[rb] = ra

    apart = {frozenset(map(str.lower, pair)) for pair in aliases["keep_apart"]}
    by_strict, by_domain = defaultdict(list), defaultdict(list)
    web = {}
    for c in companies:
        d = domain(c.get("website"))
        if d:
            by_domain[d].append(c["company"])
            web[c["company"]] = d
    for n in names:
        by_strict[strict_key(n)].append(n)
    for group in by_strict.values():
        for other in group[1:]:
            if frozenset((group[0].lower(), other.lower())) not in apart:
                union(group[0], other)
    # same website domain -> merge only when the names also share a distinctive key (units of one association share a domain)
    for group in by_domain.values():
        for i, x in enumerate(group):
            for y in group[i + 1:]:
                kx, ky = loose_keys(x), loose_keys(y)
                share = bool(kx & ky)
                if share and frozenset((x.lower(), y.lower())) not in apart:
                    union(x, y)
    lower_map = {n.lower(): n for n in names}
    for group in aliases["merge"]:
        present = [lower_map[g.lower()] for g in group if g.lower() in lower_map]
        for other in present[1:]:
            union(present[0], other)
    # candidate (unmerged) near-duplicates for manual review
    by_loose = defaultdict(set)
    for n in names:
        for k in loose_keys(n):
            by_loose[k].add(n)
    for k, grp in by_loose.items():
        roots = {find(n) for n in grp}
        if len(roots) > 1:
            report["dup_candidates"].append(f"[{k}] " + " || ".join(sorted(grp)))

    clusters = defaultdict(list)
    for c in companies:
        clusters[find(c["company"])].append(c)

    def filled(c):
        return sum(1 for k in ["website", "hq_address", "office_phone", "general_email", "cp_channel_contact", "geography",
                               "active_projects", "scale_indicators", "accessibility", "notes"] if not is_na(c.get(k)))

    merged = OrderedDict()   # root -> merged company dict
    for root, recs in clusters.items():
        recs = sorted(recs, key=lambda c: (c["_topic"].startswith("kb_"), -filled(c)))
        prim = recs[0]
        m = {k: prim.get(k) for k in prim}
        m["company"] = prim["company"]
        m["_topics"] = uniq([r["_topic"] for r in recs])
        m["_aliases"] = uniq([r["company"] for r in recs])
        m["_basis"] = basis_of(prim["_topic"])
        m["_confidence"] = s(prim.get("confidence"))
        for k in ["office_phone", "general_email", "cp_channel_contact", "website", "hq_address", "geography", "active_projects",
                  "scale_indicators", "level_rationale", "accessibility"]:
            if is_na(m.get(k)):
                for r in recs[1:]:
                    if not is_na(r.get(k)) and basis_of(r["_topic"]) == m["_basis"]:
                        m[k] = r.get(k)
                        break
        notes = uniq([s(r.get("notes")) for r in recs if not is_na(r.get("notes"))])
        m["notes"] = " | ".join(notes)
        m["sources"] = uniq(sum((src_list(r.get("sources")) for r in recs), []))
        lv = [s(r.get("level")).upper().replace("LEVEL ", "L").replace("LEVEL", "L") for r in recs]
        lv = [("Connector" if "CONNECT" in x else x.split()[0].split("/")[0].split("-")[0] if x else "") for x in lv]
        lv = [x for x in lv if x in LEVEL_ORDER]
        m["level"] = lv[0] if lv else "L4"
        if len(set(lv)) > 1:
            report["level_conflicts"].append(f"{m['company']}: {lv} -> {m['level']}")
        ov = {k.lower(): v for k, v in aliases["level_override"].items()}
        for a in m["_aliases"]:
            if a.lower() in ov:
                m["level"] = ov[a.lower()]
        merged[root] = m

    # stub companies for people whose company has no record
    for p in people:
        root = find(p["company"]) if p["company"] in parent else None
        if root and root not in merged:
            merged[root] = {"company": p["company"], "level": "Connector" if p["_topic"].startswith("assoc") else "L4",
                            "_topics": [p["_topic"]], "_aliases": [p["company"]], "sources": [], "notes": "",
                            "website": "", "hq_address": "", "geography": "", "active_projects": ""}
            report["stub_companies"].append(p["company"])

    # ---------------- QA contact checks
    checks = {}
    for pth in glob.glob(CHECKS_GLOB):
        rows, _ = load_jsonl(pth)
        for r in rows:
            key = (norm_person(r.get("person")) or "#company", strict_key(s(r.get("company"))), s(r.get("field")).lower(), re.sub(r"\s", "", s(r.get("value")).lower()))
            checks[key] = r

    def check_status(person, company, field, value):
        if is_na(value):
            return None
        key = (norm_person(person) or "#company", strict_key(company), field, re.sub(r"\s", "", s(value).lower()))
        r = checks.get(key)
        return s(r.get("status")).lower() if r else None

    # ---------------- people dedupe
    pmerged = OrderedDict()
    people.sort(key=lambda p: p["_topic"].startswith("kb_"))
    for p in people:
        root = find(p["company"])
        k = (root, norm_person(p["person"]))
        if k not in pmerged:
            q = dict(p)
            q["sources"] = src_list(p.get("sources"))
            q["_topics"] = [p["_topic"]]
            pmerged[k] = q
        else:
            q = pmerged[k]
            report["people_dupes_merged"].append(f"{p['person']} @ {p['company']}")
            same_basis = basis_of(p["_topic"]) == basis_of(q["_topic"])
            for f in ["mobile", "email", "linkedin", "designation", "role_verified", "priority", "role_category"]:
                if same_basis and is_na(q.get(f)) and not is_na(p.get(f)):
                    q[f] = p.get(f)
            if same_basis and not is_na(p.get("notes")) and s(p.get("notes")) not in s(q.get("notes")):
                q["notes"] = (s(q.get("notes")) + " | " + s(p.get("notes"))).strip(" |")
            if same_basis:
                q["sources"] = uniq(q["sources"] + src_list(p.get("sources")))
            elif basis_of(p["_topic"]) == KB_LABEL:
                q["_kb_also"] = True
            q["_topics"] = uniq(q["_topics"] + [p["_topic"]])

    # cross-link association office bearers <-> developer records
    name_index = defaultdict(list)
    for (root, pn), q in pmerged.items():
        name_index[pn].append((root, q))
    for pn, lst in name_index.items():
        if len(lst) > 1:
            for root, q in lst:
                others = [f"{merged[r]['company']} ({merged[r]['level']})" for r, _ in lst if r != root]
                q["_also"] = "Also listed under: " + "; ".join(others)

    # ---------------- build rows
    def role_cat(q):
        rc = s(q.get("role_category"))
        for r in ROLE_ORDER:
            if rc.lower().startswith(r.lower().split("/")[0]):
                return r
        return rc or "Other"

    def prio_rank(q):
        p = s(q.get("priority")).lower()
        return 0 if p.startswith("high") else 1 if p.startswith("med") else 2 if p.startswith("low") else 3

    def verification_text(q, company):
        lines = [f"Basis: {basis_of(q['_topic'])}" + (f" (confidence: {s(q.get('confidence'))})" if s(q.get("confidence")) else "")]
        for f in ["email", "mobile"]:
            st = check_status(q.get("person"), company, f, q.get(f))
            if st:
                lines.append(f"{f.capitalize()} QA re-check: {'confirmed on cited source' if st.startswith('confirm') else 'NOT re-confirmed – verify before use' if st.startswith('not') else 'source unreachable – verify before use'}")
        rv = s(q.get("role_verified"))
        if rv:
            lines.append("Role: " + rv)
        srcs = q.get("sources") or []
        if srcs:
            srcs = [x for x in srcs if x.lower().startswith("http")]
            if srcs:
                lines.append("Sources: " + " ; ".join(srcs[:6]))
        elif not is_na(q.get("email")) or not is_na(q.get("mobile")):
            lines.append("WARNING: contact listed without source URL – treat as unverified")
        return "\n".join(lines)

    people_rows = []
    have_people = set()
    for (root, pn), q in pmerged.items():
        c = merged[root]
        have_people.add(root)
        mob, em = clean_contact(q.get("mobile")), clean_contact(q.get("email"))
        pbasis = basis_of(q["_topic"])
        if pbasis == KB_LABEL:
            mob = NV if mob in (NA, NV) else mob
            em = NV if em in (NA, NV) else em
        # drop contacts that QA found absent from their cited source
        for fld in ("email", "mobile"):
            val = em if fld == "email" else mob
            st = check_status(q.get("person"), c["company"], fld, val)
            if st and st.startswith("not"):
                q.setdefault("_qa_removed", []).append(f"{fld} '{val}' could not be found on its cited source during QA and was removed")
                if fld == "email":
                    em = NA
                else:
                    mob = NA
        note = s(q.get("notes"))
        if q.get("_also"):
            note = (note + " | " + q["_also"]).strip(" |")
        if q.get("_qa_removed"):
            note = (note + " | QA: " + "; ".join(q["_qa_removed"])).strip(" |")
        people_rows.append({
            "Level": c["level"], "Builder/Developer": c["company"], "Person Name": s(q.get("person")),
            "Designation": s(q.get("designation")), "Mobile Number": mob, "Email ID": em,
            "LinkedIn": clean_link(q.get("linkedin")), "_li_search": linkedin_search(s(q.get("person")), c["company"]),
            "Company Website": s(c.get("website")) or NF,
            "Location": s(c.get("hq_address")) or NA, "Geography": s(c.get("geography")),
            "Active Projects": s(c.get("active_projects")), "Notes": note,
            "Source/Verification": verification_text(q, c["company"]),
            "Role Category": role_cat(q), "Outreach Priority": s(q.get("priority")) or "Medium",
            "Contact Availability": contact_availability(mob, em, pbasis),
            "Data Basis": pbasis,
            "Company Office Phone": clean_contact(c.get("office_phone")),
            "Company Email": clean_contact(c.get("general_email")),
            "Channel Partner Contact": clean_contact(c.get("cp_channel_contact")),
            "_sort": (LEVEL_ORDER.get(c["level"], 9), c["company"].lower(), ROLE_RANK.get(role_cat(q), 99), prio_rank(q)),
        })
    for root, c in merged.items():
        if root not in have_people:
            people_rows.append({
                "Level": c["level"], "Builder/Developer": c["company"], "Person Name": "(Promoter not publicly identified – company contact)",
                "Designation": "Company office / sales desk", "Mobile Number": NA, "Email ID": NA, "LinkedIn": NF,
                "Company Website": s(c.get("website")) or NF, "Location": s(c.get("hq_address")) or NA,
                "Geography": s(c.get("geography")), "Active Projects": s(c.get("active_projects")),
                "Notes": s(c.get("accessibility")),
                "Source/Verification": f"Basis: {c.get('_basis', WEB_LABEL)}\n" + "Sources: " + " ; ".join([x for x in (c.get("sources") or []) if x.startswith("http")][:6]),
                "Data Basis": c.get("_basis", WEB_LABEL),
                "Role Category": "Company contact", "Outreach Priority": "Low – identify promoter via MahaRERA/site visit",
                "Contact Availability": ("Not verified yet – check website/MahaRERA" if c.get("_basis") == KB_LABEL else contact_availability(None, None)) if is_na(c.get("office_phone")) and is_na(c.get("general_email")) else "Role/company email or line only",
                "Company Office Phone": clean_contact(c.get("office_phone")), "Company Email": clean_contact(c.get("general_email")),
                "Channel Partner Contact": clean_contact(c.get("cp_channel_contact")),
                "_sort": (LEVEL_ORDER.get(c["level"], 9), c["company"].lower(), 99, 3),
            })
    people_rows.sort(key=lambda r: r["_sort"])

    comp_rows = sorted(merged.values(), key=lambda c: (LEVEL_ORDER.get(c["level"], 9), c["company"].lower()))

    # ---------------- workbook
    wb = Workbook()
    FONT = "Arial"
    hdr_fill = PatternFill("solid", start_color="1F3864")
    hdr_font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
    body_font = Font(name=FONT, size=9)
    link_font = Font(name=FONT, size=9, color="0563C1", underline="single")
    input_fill = PatternFill("solid", start_color="FFF2CC")
    level_fill = {"L1": "C6E0B4", "L2": "DDEBF7", "L3": "FCE4D6", "L4": "EDEDED", "L5": "FFF2CC", "Connector": "E4DFEC"}
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap_top = Alignment(wrap_text=True, vertical="top")

    def write_table(ws, headers, rows, widths, table_name, link_cols=(), level_col=None, input_cols=()):
        for j, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=j, value=h)
            cell.font, cell.fill, cell.border = hdr_font, hdr_fill, border
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for i, r in enumerate(rows, 2):
            for j, h in enumerate(headers, 1):
                v = r.get(h, "")
                if isinstance(v, str) and len(v) > 32000:
                    v = v[:32000] + " …"
                cell = ws.cell(row=i, column=j, value=v)
                cell.font, cell.alignment, cell.border = body_font, wrap_top, border
                if h == "LinkedIn" and v == NF and r.get("_li_search"):
                    cell.value = "Not found – click to search"
                    cell.hyperlink = r["_li_search"]
                    cell.font = link_font
                elif h in link_cols and isinstance(v, str) and v.startswith("http"):
                    cell.hyperlink = v.split()[0].split(";")[0]
                    cell.font = link_font
                if h == level_col and v in level_fill:
                    cell.fill = PatternFill("solid", start_color=level_fill[v])
                    cell.font = Font(name=FONT, size=9, bold=True)
                if h in input_cols:
                    cell.fill = input_fill
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.row_dimensions[1].height = 30
        ref = f"A1:{get_column_letter(len(headers))}{max(2, len(rows) + 1)}"
        tab = Table(displayName=table_name, ref=ref)
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
        ws.add_table(tab)

    # Sheet 1: Read Me
    ws0 = wb.active
    ws0.title = "Read Me"
    ex = json.load(open(EXAMPLES_FILE)) if os.path.exists(EXAMPLES_FILE) else {}
    lines = [
        ("Mumbai / MMR Builders & Developers – Outreach Database", "title"),
        (f"Data as of {AS_OF} (compiled {BUILD_DATE}). Built for business development, channel partnerships, redevelopment and relationship building across Mumbai city, suburbs, Thane, Navi Mumbai, Raigad, Kalyan-Dombivli-Bhiwandi and Vasai-Virar-Palghar.", "p"),
        ("", "p"),
        ("How the levels are defined", "h"),
        ("L1 – National/listed majors & top-tier Mumbai groups: very large MMR sales (roughly ≥ ₹3,000 cr/yr) and land banks across several micro-markets. Access is formal – go through BD, land, redevelopment and channel-sales heads, not the chairman.", "p"),
        ("L2 – Large established Mumbai/MMR developers: roughly ₹750–3,000 cr/yr or ~10+ active MahaRERA projects; promoter-led with professional teams. Next-generation promoters and redevelopment/BD heads are the best entry points.", "p"),
        ("L3 – Mid-size regional developers: roughly ₹150–750 cr/yr or ~4–10 active projects in 1–3 micro-markets. Promoters are reachable directly; strong fit for channel mandates and JDAs.", "p"),
        ("L4 – Small/boutique developers and society-redevelopment / SRA specialists: 1–4 active projects, promoter-run and highly accessible; the main source of redevelopment deal flow.", "p"),
        ("L5 – Emerging, first-generation and single-project builders and new entrants (including JV / development-management platforms).", "p"),
        ("Connector – CREDAI-MCHI, NAREDCO and other association office bearers who can make introductions across the ecosystem.", "p"),
        ("Where sales figures are not public (most private builders), levels use proxies: number of MahaRERA-registered projects, project sizes, number of micro-markets, years active and media coverage.", "p"),
        ("", "p"),
        ("Example companies per level (full lists are in the database tabs)", "h"),
    ]
    for lvl in ["L1", "L2", "L3", "L4", "L5", "Connector"]:
        if ex.get(lvl):
            lines.append((f"{lvl}: {ex[lvl]}", "p"))
    lines += [
        ("", "p"),
        ("Tabs", "h"),
        ("Outreach Database – one row per decision-maker (the requested columns first, then role category, priority, contact availability, company lines and blank CRM tracking columns).", "p"),
        ("Company Directory – one row per developer with website, office address, published office phone/email, channel-partner contact, geography, active projects, scale indicators, level rationale and best entry point.", "p"),
        ("Connectors – association office bearers and their own developer affiliations.", "p"),
        ("Level Summary – live counts by level (formulas) of companies, contacts and contact availability.", "p"),
        ("", "p"),
        ("Contact-detail rules used", "h"),
        ("• Emails and phone numbers appear ONLY where they were seen verbatim on a public source (company website, annual report / exchange filing, MahaRERA, association directory, press release). Nothing was guessed or pattern-constructed, and data-broker 'predicted' emails were excluded.", "p"),
        (f"• '{NA}' means no public professional contact was found for that person – use the Company Office Phone / Company Email / Channel Partner Contact columns on the same row instead.", "p"),
        (f"• '{NV}' means the detail has not been checked yet (knowledge-draft rows, or pages that could not be opened in this session) – it does NOT mean the detail is unavailable.", "p"),
        ("• Data Basis – 'Web-sourced (search extracts, Sep 2026)': found by web search in this session. The session's network policy blocked opening most pages (MahaRERA, company sites, news sites), so details were read from search-result extracts of the cited pages: spot-check on the page before first use. 'Model knowledge – unverified': rows that fill in the wider developer universe from the assistant's training knowledge (to mid-2026); they carry no contact details and the roles must be confirmed on LinkedIn / the company website / MahaRERA before outreach.", "p"),
        ("• LinkedIn: where no profile was verified, the cell links to a LinkedIn people-search for that name and company.", "p"),
        ("• 'Contact Availability' classifies what is public: a direct/personal email or mobile, versus role mailboxes (info@, sales@, investor@, cs@) and board lines.", "p"),
        ("• The Source/Verification column gives the page(s) used and how the current role was confirmed; where a QA re-check was run it states whether the contact was re-confirmed on the cited page.", "p"),
        ("• Roles change often. Re-confirm the role on LinkedIn or the company website before first outreach, and treat any contact older than 12 months as stale.", "p"),
        ("", "p"),
        ("Compliance note", "h"),
        ("Use these contacts for professional, relevant B2B outreach only. India's Digital Personal Data Protection Act 2023 and its 2025 Rules apply to personal data such as individual emails and mobiles; TRAI's commercial-communication rules (DND/TCCCPR) govern unsolicited calls and SMS. Prefer email / LinkedIn introductions, honour opt-outs, and do not bulk-message mobiles.", "p"),
        ("", "p"),
        ("Suggested outreach sequence", "h"),
        ("1) Start with Connectors (CREDAI-MCHI / NAREDCO office bearers and unit presidents) for warm introductions.  2) Work L3–L4 promoters directly for channel mandates, redevelopment and JDA opportunities.  3) Approach L1–L2 through their channel-partner programmes (empanelment), then BD / land / redevelopment heads with a specific proposal.  4) Log every touch in the yellow CRM columns of the Outreach Database.", "p"),
    ]
    ws0.column_dimensions["A"].width = 150
    for i, (txt, kind) in enumerate(lines, 1):
        c = ws0.cell(row=i, column=1, value=txt)
        if kind == "title":
            c.font = Font(name=FONT, size=16, bold=True, color="1F3864")
        elif kind == "h":
            c.font = Font(name=FONT, size=11, bold=True, color="1F3864")
        else:
            c.font = Font(name=FONT, size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")

    # Sheet 2: Outreach Database
    ws1 = wb.create_sheet("Outreach Database")
    H1 = ["Level", "Builder/Developer", "Person Name", "Designation", "Mobile Number", "Email ID", "LinkedIn", "Company Website",
          "Location", "Geography", "Active Projects", "Notes", "Source/Verification", "Role Category", "Outreach Priority",
          "Contact Availability", "Data Basis", "Company Office Phone", "Company Email", "Channel Partner Contact",
          "Outreach Status", "Last Contact Date", "Next Step / Owner"]
    W1 = [9, 28, 24, 28, 18, 26, 24, 28, 32, 30, 45, 50, 60, 18, 30, 24, 22, 22, 26, 28, 16, 14, 26]
    for i, r in enumerate(people_rows, 1):
        r["Outreach Status"] = "Not contacted"
        r["Last Contact Date"] = ""
        r["Next Step / Owner"] = ""
    write_table(ws1, H1, people_rows, W1, "OutreachDB", link_cols=("LinkedIn", "Company Website"), level_col="Level",
                input_cols=("Outreach Status", "Last Contact Date", "Next Step / Owner"))
    ws1.freeze_panes = "D2"
    dv = DataValidation(type="list", formula1='"Not contacted,Contacted,Intro requested,Meeting set,In discussion,Mandate / CP empanelled,Deal in progress,Parked,Not relevant"', allow_blank=True)
    ws1.add_data_validation(dv)
    dv.add(f"U2:U{len(people_rows) + 1}")
    ws1.cell(row=1, column=21).comment = Comment("Yellow columns are for your own CRM tracking – pick a status from the dropdown.", "Database")

    # Sheet 3: Company Directory
    ws2 = wb.create_sheet("Company Directory")
    H2 = ["Level", "Builder/Developer", "Company Website", "Location", "Office Phone", "Company Email", "Channel Partner Contact",
          "Geography", "Active Projects", "Scale Indicators", "Level Rationale", "Best Entry Point / Accessibility", "Notes",
          "Contacts in Database", "Also Known As", "Data Basis", "Confidence", "Source/Verification", "Last Verified"]
    W2 = [9, 30, 28, 34, 22, 28, 28, 32, 50, 40, 40, 45, 55, 11, 28, 22, 11, 60, 14]
    crow = []
    for i, c in enumerate(comp_rows, 2):
        crow.append({
            "Level": c["level"], "Builder/Developer": c["company"], "Company Website": s(c.get("website")) or NF,
            "Location": s(c.get("hq_address")) or NA, "Office Phone": clean_contact(c.get("office_phone")),
            "Company Email": clean_contact(c.get("general_email")), "Channel Partner Contact": clean_contact(c.get("cp_channel_contact")),
            "Geography": s(c.get("geography")), "Active Projects": s(c.get("active_projects")),
            "Scale Indicators": s(c.get("scale_indicators")), "Level Rationale": s(c.get("level_rationale")),
            "Best Entry Point / Accessibility": s(c.get("accessibility")), "Notes": s(c.get("notes")),
            "Contacts in Database": f"=COUNTIFS('Outreach Database'!$B:$B,B{i},'Outreach Database'!$N:$N,\"<>Company contact\")",
            "Also Known As": "; ".join(a for a in c.get("_aliases", []) if a != c["company"]),
            "Data Basis": c.get("_basis", WEB_LABEL), "Confidence": c.get("_confidence") or ("" if c.get("_basis") != KB_LABEL else "Medium"),
            "Source/Verification": "\n".join([x for x in (c.get("sources") or []) if x.startswith("http")][:10]) or ("Model knowledge – verify on company website / MahaRERA" if c.get("_basis") == KB_LABEL else ""),
            "Last Verified": (s(c.get("last_verified")) or AS_OF) if c.get("_basis") != KB_LABEL else "Not yet verified",
        })
    write_table(ws2, H2, crow, W2, "CompanyDirectory", link_cols=("Company Website",), level_col="Level")
    ws2.freeze_panes = "C2"

    # Sheet 4: Connectors
    ws3 = wb.create_sheet("Connectors")
    conn = [r for r in people_rows if r["Level"] == "Connector"]
    H3 = ["Association", "Person Name", "Designation", "Developer Affiliation / Notes", "Mobile Number", "Email ID", "LinkedIn",
          "Association Phone", "Association Email", "Outreach Priority", "Source/Verification"]
    W3 = [30, 24, 30, 55, 18, 26, 30, 22, 26, 30, 60]
    crows = [{"Association": r["Builder/Developer"], "Person Name": r["Person Name"], "Designation": r["Designation"],
              "Developer Affiliation / Notes": r["Notes"], "Mobile Number": r["Mobile Number"], "Email ID": r["Email ID"],
              "LinkedIn": r["LinkedIn"], "Association Phone": r["Company Office Phone"], "Association Email": r["Company Email"],
              "Outreach Priority": r["Outreach Priority"], "Source/Verification": r["Source/Verification"]} for r in conn]
    write_table(ws3, H3, crows, W3, "Connectors", link_cols=("LinkedIn",))
    ws3.freeze_panes = "C2"

    # Sheet 5: Level Summary (formulas)
    ws4 = wb.create_sheet("Level Summary")
    H4 = ["Level", "Companies", "Companies (web-sourced)", "Decision-maker contacts", "Contacts (web-sourced)",
          "Contacts with public email", "Contacts with public mobile/phone", "Direct email or mobile public", "High-priority contacts"]
    for j, h in enumerate(H4, 1):
        cell = ws4.cell(row=1, column=j, value=h)
        cell.font, cell.fill, cell.border = hdr_font, hdr_fill, border
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    levels = ["L1", "L2", "L3", "L4", "L5", "Connector"]
    OD, CD = "'Outreach Database'", "'Company Directory'"
    notco = f'{OD}!$N:$N,"<>Company contact"'
    for i, lv in enumerate(levels, 2):
        vals = [lv,
                f'=COUNTIF({CD}!$A:$A,A{i})',
                f'=COUNTIFS({CD}!$A:$A,A{i},{CD}!$P:$P,"Web*")',
                f'=COUNTIFS({OD}!$A:$A,A{i},{notco})',
                f'=COUNTIFS({OD}!$A:$A,A{i},{notco},{OD}!$Q:$Q,"Web*")',
                f'=COUNTIFS({OD}!$A:$A,A{i},{notco},{OD}!$F:$F,"<>{NA}",{OD}!$F:$F,"<>{NV}")',
                f'=COUNTIFS({OD}!$A:$A,A{i},{notco},{OD}!$E:$E,"<>{NA}",{OD}!$E:$E,"<>{NV}")',
                f'=COUNTIFS({OD}!$A:$A,A{i},{notco},{OD}!$P:$P,"Direct*")+COUNTIFS({OD}!$A:$A,A{i},{notco},{OD}!$P:$P,"Mobile*")',
                f'=COUNTIFS({OD}!$A:$A,A{i},{notco},{OD}!$O:$O,"High*")']
        for j, v in enumerate(vals, 1):
            cell = ws4.cell(row=i, column=j, value=v)
            cell.font, cell.border = body_font, border
            if j == 1:
                cell.fill = PatternFill("solid", start_color=level_fill[lv])
                cell.font = Font(name=FONT, size=9, bold=True)
    tr = len(levels) + 2
    ws4.cell(row=tr, column=1, value="Total").font = Font(name=FONT, size=9, bold=True)
    for j in range(2, len(H4) + 1):
        col = get_column_letter(j)
        cell = ws4.cell(row=tr, column=j, value=f"=SUM({col}2:{col}{tr - 1})")
        cell.font, cell.border = Font(name=FONT, size=9, bold=True), border
    ws4.cell(row=tr + 2, column=1, value="Counts are live formulas over the Outreach Database and Company Directory tabs; 'Company contact' placeholder rows (no named person) are excluded from contact counts.").font = Font(name=FONT, size=9, italic=True)
    # static snapshot (for viewers that cannot calculate formulas)
    sr = tr + 4
    ws4.cell(row=sr, column=1, value=f"Snapshot at build ({BUILD_DATE}) – static values, same definitions as the live table above").font = Font(name=FONT, size=10, bold=True, color="1F3864")
    for j, h in enumerate(H4, 1):
        cell = ws4.cell(row=sr + 1, column=j, value=h)
        cell.font, cell.fill, cell.border = hdr_font, hdr_fill, border
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    real = [r for r in people_rows if r["Role Category"] != "Company contact"]
    tot = [0] * (len(H4) - 1)
    for k, lv in enumerate(levels):
        pr = [r for r in real if r["Level"] == lv]
        vals = [sum(1 for c in comp_rows if c["level"] == lv),
                sum(1 for c in comp_rows if c["level"] == lv and c.get("_basis") != KB_LABEL),
                len(pr), sum(1 for r in pr if r["Data Basis"] == WEB_LABEL),
                sum(1 for r in pr if r["Email ID"] not in (NA, NV)),
                sum(1 for r in pr if r["Mobile Number"] not in (NA, NV)),
                sum(1 for r in pr if r["Contact Availability"].startswith(("Direct", "Mobile"))),
                sum(1 for r in pr if s(r["Outreach Priority"]).lower().startswith("high"))]
        tot = [a + b for a, b in zip(tot, vals)]
        for j, v in enumerate([lv] + vals, 1):
            cell = ws4.cell(row=sr + 2 + k, column=j, value=v)
            cell.font, cell.border = body_font, border
            if j == 1:
                cell.fill = PatternFill("solid", start_color=level_fill[lv])
                cell.font = Font(name=FONT, size=9, bold=True)
    for j, v in enumerate(["Total"] + tot, 1):
        cell = ws4.cell(row=sr + 2 + len(levels), column=j, value=v)
        cell.font, cell.border = Font(name=FONT, size=9, bold=True), border
    ws4.freeze_panes = "B2"
    for j, w in enumerate([12, 12, 14, 14, 14, 16, 18, 16, 14], 1):
        ws4.column_dimensions[get_column_letter(j)].width = w

    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)

    # ---------------- build report
    stats = {
        "companies_raw": len(companies), "companies_merged": len(merged), "people_raw": len(people),
        "people_merged": len(pmerged), "rows": len(people_rows),
        "by_level_companies": {lv: sum(1 for c in merged.values() if c["level"] == lv) for lv in levels},
        "by_level_people": {lv: sum(1 for r in people_rows if r["Level"] == lv and r["Role Category"] != "Company contact") for lv in levels},
        "people_with_email": sum(1 for r in people_rows if r["Email ID"] != NA),
        "people_with_mobile": sum(1 for r in people_rows if r["Mobile Number"] != NA),
        "availability": {k: sum(1 for r in people_rows if r["Contact Availability"] == k) for k in set(r["Contact Availability"] for r in people_rows)},
        "no_source_contacts": sum(1 for r in people_rows if "WARNING" in r["Source/Verification"]),
    }
    json.dump({"stats": stats, "report": report}, open(os.path.join(SCRATCH, "build_report.json"), "w"), indent=1, ensure_ascii=False)
    print(json.dumps(stats, indent=1, ensure_ascii=False))
    for k, v in report.items():
        print(f"\n== {k} ({len(v)})")
        for x in v[:200]:
            print("  ", x)


if __name__ == "__main__":
    main()
